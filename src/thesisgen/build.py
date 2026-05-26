"""Build orchestrator: manifest → ``.docx``.

Run via ``python scripts/thesis_build.py <thesis_root>``. Reads the manifest,
auto-discovers section files per chapter, renders the front-matter, all
chapters, the bibliography, and the appendices, and writes
``build/thesis.docx``.

Manifest shape (see ``thesis_doc/manifest.yaml`` for the full example):

    front_matter: { referat, introduction }
    chapters:    [ { number, title, dir }, ... ]
    back_matter: { conclusion, appendices: [ { file, letter, title }, ... ] }
    bibliography: { source, order }
    figures: { registry }
    skip_patterns: [ "*_outline.md", ... ]
"""

from __future__ import annotations

import fnmatch
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import yaml
from docx import Document as new_document
from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from src.thesisgen import bibliography as bib
from src.thesisgen import citations as cite
from src.thesisgen import renderer
from src.thesisgen.fields import enable_field_auto_update, insert_page_number, insert_toc
from src.thesisgen.figures import FigureRegistry, load_registry
from src.thesisgen import math as math_module
from src.thesisgen.math import MathFailure
from src.thesisgen.styles import apply_gost_styles, set_cyrillic_font
from src.thesisgen.typography import normalize

APPENDIX_LETTERS = list("АБВГДЕЖИКЛМНПРСТУФХЦШЩЭЮЯ")


@dataclass
class BuildReport:
    out_path: Path
    prose_hits: list[cite.ProseHit]
    math_failures: list[MathFailure]
    n_figures: int
    n_tables: int
    n_sources: int
    n_math: int


@dataclass
class _ChapterBundle:
    number: int
    title: str
    files: list[Path]


def build_thesis(root: Path) -> BuildReport:
    root = Path(root)
    manifest = yaml.safe_load((root / "manifest.yaml").read_text(encoding="utf-8"))
    meta = yaml.safe_load((root / "meta.yaml").read_text(encoding="utf-8"))
    sources_path = root / manifest.get("bibliography", {}).get("source", "sources.yaml")
    sources = yaml.safe_load(sources_path.read_text(encoding="utf-8")) or {}
    figures_path = root / manifest.get("figures", {}).get("registry", "figures.yaml")
    figures = load_registry(figures_path, root)

    skip_patterns = manifest.get("skip_patterns", [])
    chapters = [_load_chapter(root, c, skip_patterns) for c in manifest["chapters"]]
    front = manifest.get("front_matter", {})
    back = manifest.get("back_matter", {})

    referat_text = _read(root, front.get("referat"))
    intro_text = _read(root, front.get("introduction"))
    conclusion_text = _read(root, back.get("conclusion"))
    appendix_specs = back.get("appendices", []) or []
    appendix_bodies = [
        (spec, body)
        for spec in appendix_specs
        if (body := _read(root, spec["file"])) is not None
    ]

    # Build the (filename, body) corpus for the citation map / prose scan.
    corpus: list[tuple[str, str]] = []
    if referat_text is not None:
        corpus.append((front["referat"], referat_text))
    if intro_text is not None:
        corpus.append((front["introduction"], intro_text))
    for ch in chapters:
        for f in ch.files:
            corpus.append((str(f.relative_to(root)), f.read_text(encoding="utf-8")))
    if conclusion_text is not None:
        corpus.append((back["conclusion"], conclusion_text))
    for spec, body in appendix_bodies:
        corpus.append((spec["file"], body))

    citemap = cite.build_citation_map(
        corpus, sources, manifest.get("bibliography", {}).get("order", "appearance")
    )

    # Apply citation rewrite to every source so [@key] becomes [N].
    rewritten = {filename: cite.rewrite(body, citemap) for filename, body in corpus}

    # Batch-convert every unique LaTeX formula via pandoc up-front so each
    # render call just does a cache lookup. Pandoc is the reference impl,
    # so Word will accept the OMML it produces.
    #
    # We walk each body through the same markdown-it AST the renderer uses,
    # so the captured math source matches post-CommonMark-unescaping (e.g.
    # `7\,560` in markdown becomes `7,560` here, matching what the renderer
    # will look up).
    all_math: list[str] = []
    for body in rewritten.values():
        all_math.extend(renderer.collect_math_fragments(renderer.parse(body)))
    math_module.precompute(all_math)
    prose_hits: list[cite.ProseHit] = []
    for filename, body in corpus:
        prose_hits.extend(cite.find_prose_citations(filename, body))

    doc = new_document()
    apply_gost_styles(doc)
    enable_field_auto_update(doc)
    _setup_footer(doc)
    # Title page intentionally omitted — document starts with Реферат / Содержание.

    math_failures: list[MathFailure] = []

    if referat_text is not None:
        _render_referat(doc, rewritten[front["referat"]], figures, root, math_failures, front["referat"])
    _render_toc(doc)
    if intro_text is not None:
        _render_structural_section(doc, "ВВЕДЕНИЕ", rewritten[front["introduction"]], figures, root, math_failures, front["introduction"])
    for ch in chapters:
        _render_chapter(doc, ch, rewritten, figures, root, math_failures)
    if conclusion_text is not None:
        _render_structural_section(doc, "ЗАКЛЮЧЕНИЕ", rewritten[back["conclusion"]], figures, root, math_failures, back["conclusion"])

    _render_bibliography(doc, sources, citemap)
    _render_appendices(doc, appendix_bodies, rewritten, figures, root, math_failures)

    n_figures = sum(
        1 for p in doc.paragraphs if p.style and p.style.name == "GostFigureCaption"
    )
    n_tables = len(doc.tables)
    n_sources = len(citemap)
    n_math = _count_math_in_corpus(corpus)
    _patch_referat_placeholders(doc, n_figures=n_figures, n_tables=n_tables, n_sources=n_sources)

    out_dir = root / "build"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / "thesis.docx"
    doc.save(out)
    return BuildReport(
        out_path=out,
        prose_hits=prose_hits,
        math_failures=math_failures,
        n_figures=n_figures,
        n_tables=n_tables,
        n_sources=n_sources,
        n_math=n_math,
    )


def _count_math_in_corpus(corpus: list[tuple[str, str]]) -> int:
    n = 0
    for _, body in corpus:
        n += len(renderer._MATH_INLINE_RE.findall(body))
        n += len(renderer._MATH_DISPLAY_RE.findall(body))
    return n


def _read(root: Path, name: str | None) -> str | None:
    if not name:
        return None
    p = root / "content" / name
    if not p.exists():
        return None
    return p.read_text(encoding="utf-8")


def _load_chapter(root: Path, spec: dict, skip_patterns: list[str]) -> _ChapterBundle:
    if "file" in spec:
        files = [root / "content" / str(spec["file"])]
    elif "dir" in spec:
        dir_path = root / "content" / str(spec["dir"])
        files = [
            f for f in dir_path.glob("*.md")
            if not any(fnmatch.fnmatch(f.name, pat) for pat in skip_patterns)
        ]
        files.sort(key=_natural_sort_key)
    else:
        raise ValueError(f"Chapter spec {spec!r} must declare either 'file' or 'dir'.")
    return _ChapterBundle(number=spec["number"], title=spec["title"], files=files)


def _natural_sort_key(p: Path) -> list:
    """Sort 'section_1_10' after 'section_1_2'."""
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", p.name)]


def _setup_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.first_line_indent = None
    insert_page_number(footer_p)
    if section.first_page_footer.paragraphs:
        section.first_page_footer.paragraphs[0].text = ""


def _centered(doc: Document, text: str, *, bold: bool = False, all_caps: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    run = p.add_run(text.upper() if all_caps else text)
    run.bold = bold
    set_cyrillic_font(doc.styles["Normal"])


def _render_title_page(doc: Document, meta: dict) -> None:
    for line in meta.get("organization_lines", []):
        _centered(doc, line, bold=True, all_caps=True)
    doc.add_paragraph()
    if meta.get("faculty"):
        _centered(doc, meta["faculty"])
    if meta.get("department"):
        _centered(doc, f"Кафедра: {meta['department']}")
    for _ in range(6):
        doc.add_paragraph()
    _centered(doc, meta.get("work_type", "ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА"), bold=True, all_caps=True)
    if meta.get("level"):
        _centered(doc, meta["level"])
    doc.add_paragraph()
    _centered(doc, "на тему:")
    _centered(doc, meta["title"], bold=True)
    for _ in range(4):
        doc.add_paragraph()
    if meta.get("author"):
        _right_block(doc, [f"Выполнил: {meta['author']}"])
    if meta.get("group"):
        _right_block(doc, [f"Группа: {meta['group']}"])
    if meta.get("supervisor"):
        _right_block(doc, [f"Научный руководитель: {meta['supervisor']}"])
    for _ in range(4):
        doc.add_paragraph()
    _centered(doc, f"{meta.get('city', 'Санкт-Петербург')}")
    _centered(doc, f"{meta['year']}")


def _right_block(doc: Document, lines: Iterable[str]) -> None:
    for line in lines:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = None


def _structural_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style=doc.styles["GostStructuralHeading"])
    p.paragraph_format.first_line_indent = None


def _render_referat(doc: Document, body: str, figures: FigureRegistry, root: Path, math_failures: list, current_file: str) -> None:
    _structural_heading(doc, "РЕФЕРАТ")
    ctx = renderer.RenderContext(assets_root=root, figures=figures, current_file=current_file)
    renderer.render(doc, renderer.parse(body), ctx)
    math_failures.extend(ctx.math_failures)


def _render_toc(doc: Document) -> None:
    _structural_heading(doc, "СОДЕРЖАНИЕ")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    insert_toc(p)


def _render_structural_section(doc: Document, heading: str, body: str, figures: FigureRegistry, root: Path, math_failures: list, current_file: str) -> None:
    _structural_heading(doc, heading)
    ctx = renderer.RenderContext(assets_root=root, figures=figures, current_file=current_file)
    renderer.render(doc, renderer.parse(body), ctx)
    math_failures.extend(ctx.math_failures)


def _render_chapter(
    doc: Document,
    chapter: _ChapterBundle,
    rewritten: dict[str, str],
    figures: FigureRegistry,
    root: Path,
    math_failures: list,
) -> None:
    title_p = doc.add_paragraph(f"Глава {chapter.number}. {chapter.title}", style=doc.styles["Heading 1"])
    title_p.paragraph_format.first_line_indent = None
    for f in chapter.files:
        key = str(f.relative_to(root))
        body = rewritten.get(key, f.read_text(encoding="utf-8"))
        ctx = renderer.RenderContext(assets_root=root, figures=figures, prefix=f"{chapter.number}.", current_file=key)
        renderer.render(doc, renderer.parse(body), ctx)
        math_failures.extend(ctx.math_failures)


def _render_bibliography(doc: Document, sources: dict, citemap: dict[str, int]) -> None:
    if not citemap:
        return
    _structural_heading(doc, "СПИСОК ЛИТЕРАТУРЫ")
    style = doc.styles["GostBibEntry"]
    for key, n in sorted(citemap.items(), key=lambda kv: kv[1]):
        entry = sources[key]
        line = bib.render_entry(entry)
        doc.add_paragraph(f"{n}. {normalize(line)}", style=style)


def _render_appendices(
    doc: Document,
    specs_bodies: list[tuple[dict, str]],
    rewritten: dict[str, str],
    figures: FigureRegistry,
    root: Path,
    math_failures: list,
) -> None:
    fallback_letters = iter(APPENDIX_LETTERS)
    for spec, _raw_body in specs_bodies:
        letter = spec.get("letter") or next(fallback_letters)
        _structural_heading(doc, f"ПРИЛОЖЕНИЕ {letter}")
        if spec.get("title"):
            _centered(doc, spec["title"], bold=True)
        body = rewritten.get(spec["file"], _raw_body)
        ctx = renderer.RenderContext(assets_root=root, figures=figures, prefix=f"{letter}.", current_file=spec["file"])
        renderer.render(doc, renderer.parse(body), ctx)
        math_failures.extend(ctx.math_failures)


def _patch_referat_placeholders(doc: Document, *, n_figures: int, n_tables: int, n_sources: int) -> None:
    # Pages estimated from paragraph count until Word actually paginates.
    n_pages = max(1, len(doc.paragraphs) // 25)
    replacements = {
        "{{pages}}": str(n_pages),
        "{{n_figures}}": str(n_figures),
        "{{n_tables}}": str(n_tables),
        "{{n_sources}}": str(n_sources),
    }
    for p in doc.paragraphs:
        for run in p.runs:
            for token, value in replacements.items():
                if token in run.text:
                    run.text = run.text.replace(token, value)
